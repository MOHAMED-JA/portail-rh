"""Rémunération et simulation de la masse salariale.

Saisie, import et détail nominatif : personnel RH (données confidentielles,
montants chiffrés en base, jamais inscrits au journal d'audit). Simulation :
RH (détail nominatif) et Direction générale (agrégats par direction)."""
from __future__ import annotations

import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import admin_requis, administrateur_requis, utilisateur_courant
from app.models import ROLES_RH, Employe, JournalAudit, Remuneration, StatutEmploye
from app.services import hierarchie, parametres
from app.services import remuneration as svc

router = APIRouter(prefix="/api/remuneration", tags=["Rémunération"])


class RemunerationPayload(BaseModel):
    # Bornes contrôlées par svc.controler(), comme à l'import : messages en français.
    salaire_base: float | None = None
    primes_fixes: float | None = None
    salaire_net: float | None = None
    mois_payes: float = 12
    banque: str | None = Field(default=None, max_length=80)
    rib: str | None = Field(default=None, max_length=30)


class ReglagesPayload(BaseModel):
    taux_charges: float = Field(ge=0, le=60)
    jours_par_mois: float = Field(ge=15, le=31)


class SimulationPayload(BaseModel):
    annee: int = Field(default_factory=lambda: date.today().year + 1, ge=2020, le=2100)
    augmentation_generale: float = Field(default=0, ge=0, le=50)       # % pour tous
    merite: dict[str, float] = Field(default_factory=lambda: {"1": 0, "2": 2, "3": 4, "0": 0})
    mois_effet: int = Field(default=1, ge=1, le=12)                    # mois d'application dans l'année
    enveloppe: float | None = Field(default=None, ge=0)                # budget annuel supplémentaire (DT)


def _actifs(db: Session) -> list[Employe]:
    return list(db.scalars(select(Employe).where(Employe.statut != StatutEmploye.SORTI).order_by(Employe.nom, Employe.prenom)))


def _texte(valeur: float | None) -> str | None:
    return None if valeur is None else f"{valeur:.3f}"


@router.get("", summary="Rémunérations saisies (RH)")
def lister(db: Session = Depends(get_db), utilisateur: Employe = Depends(admin_requis)):
    regles = svc.reglages(db)
    lignes = []
    for e in _actifs(db):
        r = svc.lire(db, e.id)
        lignes.append({
            "employe": {"matricule": e.matricule, "nom": e.nom, "prenom": e.prenom, "poste": e.poste,
                        "direction": e.departement.nom if e.departement else "Non affecté"},
            "remuneration": None if r is None else {
                **{k: r[k] for k in ("salaire_base", "primes_fixes", "salaire_net", "mois_payes", "banque", "rib")},
                "annuel_brut": round(svc.annuel_brut(r), 3),
                "cout_employeur": round(svc.cout_employeur(r, regles["taux_charges"]), 3)},
        })
    return {"reglages": regles, "saisies": sum(1 for l in lignes if l["remuneration"]), "effectif": len(lignes),
            "collaborateurs": lignes}


@router.put("/collaborateur/{matricule}", summary="Saisir la rémunération d'un collaborateur (RH)")
def saisir(matricule: str, payload: RemunerationPayload, db: Session = Depends(get_db),
           utilisateur: Employe = Depends(admin_requis)):
    e = db.scalar(select(Employe).where(Employe.matricule == matricule.strip().upper()))
    if not e:
        raise HTTPException(status_code=404, detail="Collaborateur introuvable.")
    rib = "".join(c for c in (payload.rib or "") if c.isdigit()) or None
    problemes = svc.controler(payload.salaire_base, payload.primes_fixes, payload.salaire_net, payload.mois_payes, rib)
    if problemes:
        raise HTTPException(status_code=422, detail="Rémunération refusée : " + " ; ".join(problemes) + ".")
    r = db.get(Remuneration, e.id) or Remuneration(employe_id=e.id)
    r.salaire_base, r.primes_fixes = _texte(payload.salaire_base), _texte(payload.primes_fixes)
    r.salaire_net, r.mois_payes = _texte(payload.salaire_net), payload.mois_payes
    r.banque = (payload.banque or "").strip() or None
    r.rib = rib
    r.modifie_le, r.modifie_par_id = datetime.utcnow(), utilisateur.id
    db.add(r)
    # Montants volontairement absents du journal : il est consultable plus largement.
    db.add(JournalAudit(acteur_id=utilisateur.id, action="remuneration_modifiee", cible=e.matricule))
    db.commit()
    return {"matricule": e.matricule, "remuneration": svc.lire(db, e.id)}


@router.get("/reglages", summary="Réglages de calcul (RH)")
def lire_reglages(db: Session = Depends(get_db), utilisateur: Employe = Depends(admin_requis)):
    return svc.reglages(db)


@router.put("/reglages", summary="Modifier les réglages de calcul (administrateur RH)")
def ecrire_reglages(payload: ReglagesPayload, db: Session = Depends(get_db),
                    utilisateur: Employe = Depends(administrateur_requis)):
    parametres.ecrire(db, "reglages_remuneration", payload.model_dump())
    db.add(JournalAudit(acteur_id=utilisateur.id, action="reglages_remuneration",
                        detail=f"Charges {payload.taux_charges} % · {payload.jours_par_mois} j/mois"))
    db.commit()
    return svc.reglages(db)


COLONNES_MODELE = ["Matricule", "Nom", "Prénom", "Direction", "Salaire de base brut mensuel (DT)",
                   "Primes fixes mensuelles (DT)", "Salaire net mensuel (DT)", "Mois payés par an", "Banque", "RIB"]


@router.get("/modele.xlsx", summary="Classeur de saisie des rémunérations (RH)")
def modele(db: Session = Depends(get_db), utilisateur: Employe = Depends(admin_requis)):
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Rémunérations"
    feuille.append(COLONNES_MODELE)
    for e in _actifs(db):
        r = svc.lire(db, e.id) or {}
        feuille.append([e.matricule, e.nom, e.prenom, e.departement.nom if e.departement else "",
                        r.get("salaire_base"), r.get("primes_fixes"), r.get("salaire_net"), r.get("mois_payes", 12),
                        r.get("banque"), r.get("rib")])
    for colonne, largeur in zip("ABCDEFGHIJ", (12, 18, 18, 30, 18, 18, 18, 12, 18, 24)):
        feuille.column_dimensions[colonne].width = largeur
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return StreamingResponse(tampon, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="RH-remunerations.xlsx"'})


def _cellule(valeur) -> float | None:
    if valeur in (None, ""):
        return None
    try:
        return float(str(valeur).replace(",", ".").replace(" ", "").replace(" ", ""))
    except ValueError:
        raise ValueError(f"« {valeur} » n'est pas un montant") from None


@router.post("/import", summary="Importer le classeur des rémunérations (RH)")
def importer(fichier: UploadFile = File(...), appliquer: bool = False, db: Session = Depends(get_db),
             utilisateur: Employe = Depends(admin_requis)):
    """Sans ``appliquer`` : simulation (erreurs et nombre de lignes). Avec : écriture."""
    try:
        feuille = load_workbook(io.BytesIO(fichier.file.read()), data_only=True).active
    except Exception:
        raise HTTPException(status_code=422, detail="Classeur Excel illisible (.xlsx attendu).") from None
    erreurs, pretes = [], []
    for numero, ligne in enumerate(feuille.iter_rows(min_row=2, values_only=True), start=2):
        if not ligne or not ligne[0]:
            continue
        matricule = str(ligne[0]).strip().upper()
        e = db.scalar(select(Employe).where(Employe.matricule == matricule))
        if not e:
            erreurs.append(f"Ligne {numero} : matricule {matricule} inconnu.")
            continue
        try:
            base, primes, net = (_cellule(ligne[i]) if len(ligne) > i else None for i in (4, 5, 6))
            mois = _cellule(ligne[7]) if len(ligne) > 7 else None
        except ValueError as souci:
            erreurs.append(f"Ligne {numero} ({matricule}) : {souci}.")
            continue
        if base is None:
            continue   # ligne laissée vide : rien à importer
        rib = "".join(c for c in str(ligne[9] or "") if c.isdigit()) if len(ligne) > 9 else ""
        # Mêmes règles que la saisie individuelle, avant toute écriture.
        problemes = svc.controler(base, primes, net, mois, rib or None)
        if problemes:
            erreurs.append(f"Ligne {numero} ({matricule}) : " + " ; ".join(problemes) + ".")
            continue
        pretes.append((e, base, primes, net, mois or 12, str(ligne[8]).strip() if len(ligne) > 8 and ligne[8] else None, rib or None))
    if appliquer and not erreurs:
        for e, base, primes, net, mois, banque, rib in pretes:
            r = db.get(Remuneration, e.id) or Remuneration(employe_id=e.id)
            r.salaire_base, r.primes_fixes, r.salaire_net = _texte(base), _texte(primes), _texte(net)
            r.mois_payes, r.banque, r.rib = mois, banque, rib
            r.modifie_le, r.modifie_par_id = datetime.utcnow(), utilisateur.id
            db.add(r)
        db.add(JournalAudit(acteur_id=utilisateur.id, action="remunerations_importees", detail=f"{len(pretes)} ligne(s)"))
        db.commit()
    return {"lignes": len(pretes), "erreurs": erreurs, "applique": bool(appliquer and not erreurs)}


@router.post("/simulation", summary="Simuler les augmentations et la masse salariale (RH ; DG en agrégé)")
def simuler(payload: SimulationPayload, db: Session = Depends(get_db), utilisateur: Employe = Depends(utilisateur_courant)):
    if not hierarchie.voit_tout(utilisateur):
        raise HTTPException(status_code=403, detail="Réservé à la RH et à la Direction générale.")
    rh = utilisateur.role in ROLES_RH
    regles = svc.reglages(db)
    charges = 1 + regles["taux_charges"] / 100
    mois_restants = 13 - payload.mois_effet
    par_direction: dict[str, dict] = {}
    lignes, non_valorises = [], 0
    for e in _actifs(db):
        r = svc.lire(db, e.id)
        if r is None:
            non_valorises += 1
            continue
        niveau = svc.niveau_performance(db, e.id, payload.annee - 1)
        taux = payload.augmentation_generale + payload.merite.get(str(niveau or 0), 0)
        mensuel = svc.mensuel_brut(r)
        # L'augmentation porte sur le salaire de base ; les primes fixes restent inchangées.
        hausse_mensuelle = r["salaire_base"] * taux / 100
        actuel_annuel = mensuel * r["mois_payes"] * charges
        cout_annee = hausse_mensuelle * r["mois_payes"] * mois_restants / 12 * charges
        cout_plein = hausse_mensuelle * r["mois_payes"] * charges
        direction = e.departement.nom if e.departement else "Non affecté"
        a = par_direction.setdefault(direction, {"direction": direction, "effectif": 0, "masse_actuelle": 0.0,
                                                 "cout_annee": 0.0, "cout_annee_pleine": 0.0})
        a["effectif"] += 1
        a["masse_actuelle"] += actuel_annuel
        a["cout_annee"] += cout_annee
        a["cout_annee_pleine"] += cout_plein
        if rh:
            lignes.append({"employe": {"matricule": e.matricule, "nom": e.nom, "prenom": e.prenom, "direction": direction},
                           "performance": niveau, "taux": round(taux, 2),
                           "salaire_base": r["salaire_base"], "nouveau_salaire_base": round(r["salaire_base"] + hausse_mensuelle, 3),
                           "cout_annee": round(cout_annee, 3), "cout_annee_pleine": round(cout_plein, 3)})
    for a in par_direction.values():
        for cle in ("masse_actuelle", "cout_annee", "cout_annee_pleine"):
            a[cle] = round(a[cle], 3)
    masse = sum(a["masse_actuelle"] for a in par_direction.values())
    cout_annee = sum(a["cout_annee"] for a in par_direction.values())
    cout_plein = sum(a["cout_annee_pleine"] for a in par_direction.values())
    return {
        "parametres": payload.model_dump(), "reglages": regles, "non_valorises": non_valorises,
        "masse_actuelle": round(masse, 3), "cout_annee": round(cout_annee, 3), "cout_annee_pleine": round(cout_plein, 3),
        "masse_apres": round(masse + cout_plein, 3),
        "evolution_pourcent": round(100 * cout_plein / masse, 2) if masse else None,
        "ecart_enveloppe": None if payload.enveloppe is None else round(payload.enveloppe - cout_annee, 3),
        "par_direction": sorted(par_direction.values(), key=lambda a: -a["masse_actuelle"]),
        "collaborateurs": sorted(lignes, key=lambda l: -l["cout_annee_pleine"]) if rh else None,
    }
