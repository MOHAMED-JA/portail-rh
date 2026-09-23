"""Import Excel contrôlé des affectations et des dossiers administratifs.

Le classeur produit par :func:`generer_modele` contient les valeurs actuelles.
Une cellule laissée vide n'entraîne aucune modification ; le mot ``EFFACER``
permet de retirer explicitement une valeur existante. L'analyse est sans écriture et
l'application est atomique : une seule erreur refuse tout le classeur.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Departement,
    DossierEmploye,
    Employe,
    EvenementCarriere,
    JournalAudit,
    StatutEmploye,
)
from app.services.hierarchie import LIBELLES_NIVEAUX

MARQUE_EFFACER = "EFFACER"
TAILLE_MAX = 10 * 1024 * 1024

ENTETES_AFFECTATIONS = [
    "Matricule", "Nom", "Prénom", "Code structure", "Matricule supérieur",
    "Niveau hiérarchique", "Poste",
]
ENTETES_DOSSIERS = [
    "Matricule", "Nom", "Prénom", "Date d'entrée", "Date de naissance",
    "Catégorie", "Grade", "Échelon", "Type de contrat", "Fin du contrat",
    "Fin de période d'essai", "Diplômes", "Contact d'urgence", "Lien",
    "Téléphone", "Périodicité visite (mois)",
]

CHAMPS_DOSSIER = {
    "Date de naissance": "date_naissance",
    "Catégorie": "categorie",
    "Grade": "grade",
    "Échelon": "echelon",
    "Type de contrat": "type_contrat",
    "Fin du contrat": "date_fin_contrat",
    "Fin de période d'essai": "date_fin_essai",
    "Diplômes": "diplomes",
    "Contact d'urgence": "contact_nom",
    "Lien": "contact_lien",
    "Téléphone": "contact_telephone",
    "Périodicité visite (mois)": "visite_periodicite_mois",
}
CHAMPS_DATES = {"Date d'entrée", "Date de naissance", "Fin du contrat", "Fin de période d'essai"}
LIMITES_TEXTE = {
    "categorie": 60, "grade": 60, "echelon": 20, "type_contrat": 30,
    "diplomes": 4000, "contact_nom": 120, "contact_lien": 60,
    "contact_telephone": 30,
}


def _normaliser(valeur: object) -> str:
    texte = unicodedata.normalize("NFKD", str(valeur or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texte.strip().lower())


def _texte(valeur: object) -> str | None:
    if valeur is None:
        return None
    texte = str(valeur).strip()
    return texte or None


def _valeur_modifiable(valeur: object) -> tuple[bool, object | None]:
    """(renseigné, valeur). Une cellule vide ne demande aucune modification."""
    texte = _texte(valeur)
    if texte is None:
        return False, None
    if _normaliser(texte) == _normaliser(MARQUE_EFFACER):
        return True, None
    return True, valeur


def _date(valeur: object, feuille: str, ligne: int, colonne: str, erreurs: list[str]) -> date | None:
    renseigne, valeur = _valeur_modifiable(valeur)
    if not renseigne or valeur is None:
        return None
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    texte = str(valeur).strip()
    for format_date in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texte, format_date).date()
        except ValueError:
            continue
    erreurs.append(f"{feuille} ligne {ligne} — {colonne} : date invalide « {texte} » (JJ/MM/AAAA attendu).")
    return None


def _nom(employe: Employe) -> str:
    return f"{employe.prenom} {employe.nom}"


def _style_feuille(feuille, largeurs: list[int], lignes: int) -> None:
    rouge, marine, gris, jaune = "DF271C", "072241", "E9EEF4", "FFF2CC"
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = f"A1:{get_column_letter(len(largeurs))}{max(2, lignes)}"
    feuille.row_dimensions[1].height = 30
    for cellule in feuille[1]:
        cellule.fill = PatternFill("solid", fgColor=marine)
        cellule.font = Font(color="FFFFFF", bold=True)
        cellule.alignment = Alignment(vertical="center", wrap_text=True)
    for index, largeur in enumerate(largeurs, 1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur
    for ligne in range(2, lignes + 1):
        for colonne in range(1, 4):
            feuille.cell(ligne, colonne).fill = PatternFill("solid", fgColor=gris)
        for colonne in range(4, len(largeurs) + 1):
            feuille.cell(ligne, colonne).fill = PatternFill("solid", fgColor=jaune)
        feuille.row_dimensions[ligne].height = 22
    feuille.conditional_formatting.add(
        f"A2:A{max(2, lignes)}",
        FormulaRule(formula=["LEN($A2)=0"], fill=PatternFill("solid", fgColor=rouge)),
    )


def _liste_validation(feuille, formule: str, plage: str) -> None:
    validation = DataValidation(type="list", formula1=formule, allow_blank=True)
    validation.error = "Choisissez une valeur de la liste."
    validation.errorTitle = "Valeur inconnue"
    validation.prompt = f"Saisissez {MARQUE_EFFACER} pour retirer une valeur existante."
    validation.promptTitle = "Valeur contrôlée"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    feuille.add_data_validation(validation)
    validation.add(plage)


def generer_modele(db: Session) -> io.BytesIO:
    """Génère un classeur prérempli avec tous les collaborateurs actifs."""
    employes = list(db.scalars(select(Employe).where(Employe.statut != StatutEmploye.SORTI)
                               .order_by(Employe.nom, Employe.prenom)))
    departements = list(db.scalars(select(Departement).order_by(Departement.code)))
    dossiers = {d.employe_id: d for d in db.scalars(select(DossierEmploye))}

    classeur = Workbook()
    instructions = classeur.active
    instructions.title = "Instructions"
    instructions.sheet_view.showGridLines = False
    instructions.column_dimensions["A"].width = 25
    instructions.column_dimensions["B"].width = 105
    instructions["A1"], instructions["B1"] = "IMPORT RH EN MASSE", "Affectations et dossiers administratifs"
    for cellule in instructions[1]:
        cellule.fill = PatternFill("solid", fgColor="072241")
        cellule.font = Font(color="FFFFFF", bold=True, size=14)
    consignes = [
        ("1. Identifier", "Ne modifiez jamais Matricule, Nom ou Prénom. Ils sécurisent la correspondance avec le bon compte."),
        ("2. Renseigner", "Complétez seulement les cellules jaunes. Les valeurs existantes sont préremplies."),
        ("3. Effacer", f"Écrivez {MARQUE_EFFACER} pour retirer explicitement une structure, un supérieur ou une donnée facultative."),
        ("4. Dates", "Utilisez de vraies dates Excel ou le format JJ/MM/AAAA."),
        ("5. Contrôler", "Le portail affiche une prévisualisation et refuse tout le fichier si une seule ligne est invalide."),
        ("6. Appliquer", "Une sauvegarde est créée automatiquement avant toute écriture. Chaque modification est tracée dans le journal d’audit."),
    ]
    for ligne, (titre, detail) in enumerate(consignes, 3):
        instructions.cell(ligne, 1, titre).font = Font(bold=True, color="DF271C")
        instructions.cell(ligne, 2, detail).alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[ligne].height = 35

    references = classeur.create_sheet("Références")
    references.append(["Structures", "Matricules", "Niveaux", "Types de contrat"])
    max_ref = max(len(departements), len(employes), len(LIBELLES_NIVEAUX), 5)
    contrats = ["CDI", "CDD", "Stage", "SIVP", "Contractuel"]
    for index in range(max_ref):
        references.append([
            departements[index].code if index < len(departements) else None,
            employes[index].matricule if index < len(employes) else None,
            list(LIBELLES_NIVEAUX)[index] if index < len(LIBELLES_NIVEAUX) else None,
            contrats[index] if index < len(contrats) else None,
        ])
    references.sheet_state = "hidden"

    affectations = classeur.create_sheet("Affectations")
    affectations.append(ENTETES_AFFECTATIONS)
    for e in employes:
        affectations.append([
            e.matricule, e.nom, e.prenom, e.departement.code if e.departement else None,
            e.validateur.matricule if e.validateur else None, e.niveau or "collaborateur", e.poste,
        ])
    _style_feuille(affectations, [16, 22, 22, 20, 22, 24, 40], len(employes) + 1)
    _liste_validation(affectations, f"'Références'!$A$2:$A${len(departements)+1}", f"D2:D{len(employes)+1}")
    _liste_validation(affectations, f"'Références'!$B$2:$B${len(employes)+1}", f"E2:E{len(employes)+1}")
    _liste_validation(affectations, f"'Références'!$C$2:$C${len(LIBELLES_NIVEAUX)+1}", f"F2:F{len(employes)+1}")

    feuille_dossiers = classeur.create_sheet("Dossiers")
    feuille_dossiers.append(ENTETES_DOSSIERS)
    for e in employes:
        d = dossiers.get(e.id)
        valeur = lambda champ: getattr(d, champ) if d else None
        feuille_dossiers.append([
            e.matricule, e.nom, e.prenom, e.date_entree, valeur("date_naissance"), valeur("categorie"),
            valeur("grade"), valeur("echelon"), valeur("type_contrat"), valeur("date_fin_contrat"),
            valeur("date_fin_essai"), valeur("diplomes"), valeur("contact_nom"), valeur("contact_lien"),
            valeur("contact_telephone"), valeur("visite_periodicite_mois") or 12,
        ])
    _style_feuille(feuille_dossiers, [16, 22, 22, 16, 18, 20, 20, 14, 18, 16, 22, 34, 24, 18, 20, 22], len(employes) + 1)
    for colonne in (4, 5, 10, 11):
        for ligne in range(2, len(employes) + 2):
            feuille_dossiers.cell(ligne, colonne).number_format = "dd/mm/yyyy"
    _liste_validation(feuille_dossiers, "'Références'!$D$2:$D$6", f"I2:I{len(employes)+1}")

    classeur.active = 0
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


@dataclass
class AnalyseImport:
    erreurs: list[str] = field(default_factory=list)
    modifications: dict[int, dict] = field(default_factory=dict)
    lignes_lues: int = 0

    def rapport(self) -> dict:
        apercu = []
        champs = 0
        for modification in self.modifications.values():
            changes = [*modification["employe"], *modification["dossier"]]
            if changes:
                champs += len(changes)
                apercu.append({
                    "matricule": modification["matricule"],
                    "nom": modification["nom"],
                    "champs": changes,
                })
        return {
            "valide": not self.erreurs,
            "lignes_lues": self.lignes_lues,
            "collaborateurs_modifies": len(apercu),
            "champs_modifies": champs,
            "erreurs": self.erreurs,
            "apercu": apercu,
        }


def _entetes(feuille, attendues: list[str], erreurs: list[str]) -> dict[str, int]:
    presentes = {_normaliser(c.value): index for index, c in enumerate(feuille[1], 1) if c.value is not None}
    absentes = [h for h in attendues if _normaliser(h) not in presentes]
    if absentes:
        erreurs.append(f"{feuille.title} — colonne(s) absente(s) : {', '.join(absentes)}.")
    return {h: presentes.get(_normaliser(h), 0) for h in attendues}


def _cellule(feuille, ligne: int, colonnes: dict[str, int], nom: str):
    return feuille.cell(ligne, colonnes[nom]).value if colonnes.get(nom) else None


def analyser(db: Session, contenu: bytes) -> AnalyseImport:
    analyse = AnalyseImport()
    if not contenu:
        analyse.erreurs.append("Le fichier Excel est vide.")
        return analyse
    if len(contenu) > TAILLE_MAX:
        analyse.erreurs.append("Le fichier dépasse la taille maximale de 10 Mo.")
        return analyse
    try:
        classeur = load_workbook(io.BytesIO(contenu), data_only=True, read_only=True)
    except Exception as erreur:  # noqa: BLE001
        analyse.erreurs.append(f"Classeur illisible : {type(erreur).__name__}.")
        return analyse
    for nom in ("Affectations", "Dossiers"):
        if nom not in classeur.sheetnames:
            analyse.erreurs.append(f"Feuille obligatoire absente : {nom}.")
    if analyse.erreurs:
        return analyse

    employes = list(db.scalars(select(Employe)))
    par_matricule = {e.matricule.upper(): e for e in employes}
    matricule_par_id = {e.id: e.matricule for e in employes}
    structures = {d.code.upper(): d for d in db.scalars(select(Departement))}
    dossiers = {d.employe_id: d for d in db.scalars(select(DossierEmploye))}

    def modification(e: Employe) -> dict:
        return analyse.modifications.setdefault(e.id, {
            "matricule": e.matricule, "nom": _nom(e), "employe": {}, "dossier": {},
        })

    # Affectations et hiérarchie.
    feuille = classeur["Affectations"]
    colonnes = _entetes(feuille, ENTETES_AFFECTATIONS, analyse.erreurs)
    doublons: set[str] = set()
    for ligne in range(2, feuille.max_row + 1):
        matricule = (_texte(_cellule(feuille, ligne, colonnes, "Matricule")) or "").upper()
        if not matricule:
            continue
        analyse.lignes_lues += 1
        if matricule in doublons:
            analyse.erreurs.append(f"Affectations ligne {ligne} — matricule {matricule} présent plusieurs fois.")
            continue
        doublons.add(matricule)
        e = par_matricule.get(matricule)
        if not e:
            analyse.erreurs.append(f"Affectations ligne {ligne} — matricule inconnu : {matricule}.")
            continue
        nom, prenom = _texte(_cellule(feuille, ligne, colonnes, "Nom")), _texte(_cellule(feuille, ligne, colonnes, "Prénom"))
        if _normaliser(nom) != _normaliser(e.nom) or _normaliser(prenom) != _normaliser(e.prenom):
            analyse.erreurs.append(f"Affectations ligne {ligne} — identité incohérente pour {matricule}.")
            continue
        m = modification(e)
        renseigne, valeur = _valeur_modifiable(_cellule(feuille, ligne, colonnes, "Code structure"))
        if renseigne:
            code = str(valeur).strip().upper() if valeur is not None else None
            departement = structures.get(code) if code else None
            if code and not departement:
                analyse.erreurs.append(f"Affectations ligne {ligne} — structure inconnue : {code}.")
            elif (departement.id if departement else None) != e.departement_id:
                m["employe"]["departement_id"] = departement.id if departement else None
        renseigne, valeur = _valeur_modifiable(_cellule(feuille, ligne, colonnes, "Matricule supérieur"))
        if renseigne:
            superieur_matricule = str(valeur).strip().upper() if valeur is not None else None
            superieur = par_matricule.get(superieur_matricule) if superieur_matricule else None
            if superieur_matricule and not superieur:
                analyse.erreurs.append(f"Affectations ligne {ligne} — supérieur inconnu : {superieur_matricule}.")
            elif superieur and superieur.statut == StatutEmploye.SORTI:
                analyse.erreurs.append(f"Affectations ligne {ligne} — le supérieur {superieur_matricule} est sorti.")
            elif superieur and superieur.id == e.id:
                analyse.erreurs.append(f"Affectations ligne {ligne} — un collaborateur ne peut pas être son propre supérieur.")
            elif (superieur.id if superieur else None) != e.validateur_id:
                m["employe"]["validateur_id"] = superieur.id if superieur else None
        renseigne, valeur = _valeur_modifiable(_cellule(feuille, ligne, colonnes, "Niveau hiérarchique"))
        if renseigne:
            niveau = str(valeur).strip().lower() if valeur is not None else ""
            if niveau not in LIBELLES_NIVEAUX:
                analyse.erreurs.append(f"Affectations ligne {ligne} — niveau hiérarchique inconnu : {valeur}.")
            elif niveau != (e.niveau or "collaborateur"):
                m["employe"]["niveau"] = niveau
        renseigne, valeur = _valeur_modifiable(_cellule(feuille, ligne, colonnes, "Poste"))
        if renseigne:
            poste = _texte(valeur)
            if not poste:
                analyse.erreurs.append(f"Affectations ligne {ligne} — le poste ne peut pas être effacé.")
            elif len(poste) > 120:
                analyse.erreurs.append(f"Affectations ligne {ligne} — poste limité à 120 caractères.")
            elif poste != e.poste:
                m["employe"]["poste"] = poste

    # Validation globale des cycles, y compris lorsque plusieurs lignes changent ensemble.
    parents = {e.id: analyse.modifications.get(e.id, {}).get("employe", {}).get("validateur_id", e.validateur_id)
               for e in employes}
    cycles_signales: set[int] = set()
    for depart in parents:
        chemin, courant = set(), depart
        while courant is not None and courant in parents:
            if courant in chemin:
                if depart not in cycles_signales:
                    analyse.erreurs.append(
                        f"Affectations — les rattachements prévus créent une boucle autour de "
                        f"{matricule_par_id.get(depart, 'un collaborateur')}."
                    )
                    cycles_signales.add(depart)
                break
            chemin.add(courant)
            courant = parents.get(courant)

    # Dossiers administratifs.
    feuille = classeur["Dossiers"]
    colonnes = _entetes(feuille, ENTETES_DOSSIERS, analyse.erreurs)
    doublons.clear()
    for ligne in range(2, feuille.max_row + 1):
        matricule = (_texte(_cellule(feuille, ligne, colonnes, "Matricule")) or "").upper()
        if not matricule:
            continue
        analyse.lignes_lues += 1
        if matricule in doublons:
            analyse.erreurs.append(f"Dossiers ligne {ligne} — matricule {matricule} présent plusieurs fois.")
            continue
        doublons.add(matricule)
        e = par_matricule.get(matricule)
        if not e:
            analyse.erreurs.append(f"Dossiers ligne {ligne} — matricule inconnu : {matricule}.")
            continue
        nom, prenom = _texte(_cellule(feuille, ligne, colonnes, "Nom")), _texte(_cellule(feuille, ligne, colonnes, "Prénom"))
        if _normaliser(nom) != _normaliser(e.nom) or _normaliser(prenom) != _normaliser(e.prenom):
            analyse.erreurs.append(f"Dossiers ligne {ligne} — identité incohérente pour {matricule}.")
            continue
        d = dossiers.get(e.id)
        m = modification(e)
        for entete in ENTETES_DOSSIERS[3:]:
            brute = _cellule(feuille, ligne, colonnes, entete)
            renseigne, valeur = _valeur_modifiable(brute)
            if not renseigne:
                continue
            champ = "date_entree" if entete == "Date d'entrée" else CHAMPS_DOSSIER[entete]
            actuelle = getattr(e, champ) if champ == "date_entree" else (getattr(d, champ) if d else (12 if champ == "visite_periodicite_mois" else None))
            if entete in CHAMPS_DATES:
                nouvelle = _date(brute, "Dossiers", ligne, entete, analyse.erreurs)
            elif champ == "visite_periodicite_mois":
                if valeur is None:
                    nouvelle = 12
                else:
                    try:
                        nouvelle = int(valeur)
                    except (TypeError, ValueError):
                        nouvelle = 0
                    if not 1 <= nouvelle <= 60:
                        analyse.erreurs.append(f"Dossiers ligne {ligne} — périodicité comprise entre 1 et 60 mois.")
                        continue
            else:
                nouvelle = _texte(valeur)
                if nouvelle is not None and len(nouvelle) > LIMITES_TEXTE[champ]:
                    analyse.erreurs.append(f"Dossiers ligne {ligne} — {entete} dépasse {LIMITES_TEXTE[champ]} caractères.")
                    continue
            if entete == "Date de naissance" and nouvelle and nouvelle > date.today():
                analyse.erreurs.append(f"Dossiers ligne {ligne} — la date de naissance est dans le futur.")
                continue
            if nouvelle != actuelle:
                cible = m["employe"] if champ == "date_entree" else m["dossier"]
                cible[champ] = nouvelle

    return analyse


def appliquer(db: Session, analyse: AnalyseImport, acteur: Employe) -> dict:
    """Applique une analyse valide dans la transaction de l'appelant."""
    if analyse.erreurs:
        raise ValueError("Le classeur contient des erreurs.")
    changements = 0
    for employe_id, modification in analyse.modifications.items():
        if not modification["employe"] and not modification["dossier"]:
            continue
        employe = db.get(Employe, employe_id)
        avant_departement = employe.departement.nom if employe.departement else None
        avant_poste = employe.poste
        for champ, valeur in modification["employe"].items():
            setattr(employe, champ, valeur)
        if modification["dossier"]:
            dossier = db.get(DossierEmploye, employe.id)
            if dossier is None:
                dossier = DossierEmploye(employe_id=employe.id)
                db.add(dossier)
            for champ, valeur in modification["dossier"].items():
                setattr(dossier, champ, valeur)
        db.flush()
        apres_departement = employe.departement.nom if employe.departement else None
        if avant_departement != apres_departement:
            db.add(EvenementCarriere(employe_id=employe.id, date_effet=date.today(), type_evenement="mutation",
                                     avant=avant_departement, apres=apres_departement, saisi_par_id=acteur.id,
                                     commentaire="Import Excel RH en masse"))
        if avant_poste != employe.poste:
            db.add(EvenementCarriere(employe_id=employe.id, date_effet=date.today(), type_evenement="changement_poste",
                                     avant=avant_poste, apres=employe.poste, saisi_par_id=acteur.id,
                                     commentaire="Import Excel RH en masse"))
        detail = {"employe": list(modification["employe"]), "dossier": list(modification["dossier"])}
        db.add(JournalAudit(acteur_id=acteur.id, action="import_rh_collaborateur", cible=employe.matricule,
                            detail=json.dumps(detail, ensure_ascii=False)))
        changements += 1
    db.add(JournalAudit(acteur_id=acteur.id, action="import_rh_masse", cible=str(changements),
                        detail=f"{analyse.lignes_lues} ligne(s) lue(s)"))
    return {**analyse.rapport(), "appliques": changements}
