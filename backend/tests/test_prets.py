"""Avances sur salaire et prêts sociaux."""
from datetime import date

from openpyxl import load_workbook
import io

from app.routers.prets import echeancier, mois_suivant, retenues_du_mois


def demander(client, entetes, type_pret="pret_social", montant=1200, n=12, qui="100259"):
    return client.post("/api/prets", headers=entetes(qui), json={"type_pret": type_pret, "montant": montant,
                                                                 "nb_mensualites": n, "motif": "Frais de scolarité"})


def test_echeancier_sans_interets():
    lignes = echeancier(1000, 3, 0, date(2026, 10, 1))
    assert [l["montant"] for l in lignes] == [333.333, 333.333, 333.334]
    assert [l["mois"] for l in lignes] == [date(2026, 10, 1), date(2026, 11, 1), date(2026, 12, 1)]
    assert round(sum(l["capital"] for l in lignes), 3) == 1000


def test_echeancier_avec_interets():
    lignes = echeancier(12000, 12, 6, date(2026, 12, 1))
    assert round(sum(l["capital"] for l in lignes), 3) == 12000       # tout le capital est remboursé
    assert lignes[0]["montant"] == 1032.797                            # mensualité constante à 6 %/an
    assert lignes[0]["interets"] == 60.0 and lignes[-1]["interets"] < lignes[0]["interets"]
    assert lignes[-1]["mois"] == date(2027, 11, 1)


def test_mois_suivant():
    assert mois_suivant(date(2026, 12, 15)) == date(2027, 1, 1)
    assert mois_suivant(date(2026, 1, 31), 13) == date(2027, 2, 1)


def test_plafonds(client, entetes):
    assert demander(client, entetes, "avance", 5000, 2).status_code == 422          # plafond 1 000 DT
    assert demander(client, entetes, "avance", 500, 6).status_code == 422           # 3 mensualités au plus
    assert demander(client, entetes, "inexistant", 500, 2).status_code == 422


def test_une_seule_demande_en_cours_par_type(client, entetes):
    assert demander(client, entetes).status_code == 201
    assert demander(client, entetes).status_code == 409
    assert demander(client, entetes, "avance", 300, 2).status_code == 201           # autre type : possible


def test_confidentialite(client, entetes):
    p = demander(client, entetes).json()
    assert client.get(f"/api/prets/{p['id']}", headers=entetes("100259")).status_code == 200
    assert client.get(f"/api/prets/{p['id']}", headers=entetes("ADMINRH")).status_code == 200
    assert client.get(f"/api/prets/{p['id']}", headers=entetes("100130")).status_code == 403   # supérieur
    assert client.get(f"/api/prets/{p['id']}", headers=entetes("100281")).status_code == 403   # collègue
    assert client.get("/api/prets", headers=entetes("100130")).status_code == 403


def test_accord_echeancier_et_notification(client, entetes):
    p = demander(client, entetes).json()
    assert client.post(f"/api/prets/{p['id']}/decision", headers=entetes("100130"), json={"accorde": True}).status_code == 403
    r = client.post(f"/api/prets/{p['id']}/decision", headers=entetes("ADMINRH"), json={"accorde": True, "montant": 1000, "nb_mensualites": 10})
    assert r.status_code == 200
    d = r.json()
    assert d["statut"] == "accorde" and len(d["echeances"]) == 10 and d["mensualite"] == 100
    assert d["premiere_echeance"] == str(mois_suivant(date.today()))
    assert d["reste_du"] == 1000
    notes = client.get("/api/notifications", headers=entetes("100259")).json()
    assert any("accordé" in n["titre"] for n in notes)
    assert client.post(f"/api/prets/{p['id']}/decision", headers=entetes("ADMINRH"), json={"accorde": True}).status_code == 409


def test_refus_avec_motif_obligatoire(client, entetes):
    p = demander(client, entetes).json()
    assert client.post(f"/api/prets/{p['id']}/decision", headers=entetes("ADMINRH"), json={"accorde": False}).status_code == 422
    r = client.post(f"/api/prets/{p['id']}/decision", headers=entetes("ADMINRH"), json={"accorde": False, "commentaire": "Budget épuisé"})
    assert r.json()["statut"] == "refuse"


def test_report_et_solde_anticipe(client, entetes):
    p = demander(client, entetes, montant=600, n=6).json()
    client.post(f"/api/prets/{p['id']}/decision", headers=entetes("ADMINRH"), json={"accorde": True})
    r = client.post(f"/api/prets/{p['id']}/reporter/1", headers=entetes("ADMINRH")).json()
    prevues = [e for e in r["echeances"] if e["statut"] == "prevue"]
    assert len(prevues) == 6 and r["echeances"][0]["statut"] == "reportee"
    assert prevues[-1]["mois"] == str(mois_suivant(date.today(), 7))
    r = client.post(f"/api/prets/{p['id']}/solder", headers=entetes("ADMINRH")).json()
    assert r["statut"] == "solde" and r["reste_du"] == 0


def test_retenues_dans_l_export_paie(client, entetes, db):
    p = demander(client, entetes, montant=1200, n=12).json()
    client.post(f"/api/prets/{p['id']}/decision", headers=entetes("ADMINRH"), json={"accorde": True})
    mois = mois_suivant(date.today())
    from app.models import Employe
    julien = db.query(Employe).filter_by(matricule="100259").one()
    assert retenues_du_mois(db, julien.id, mois) == 100
    r = client.get(f"/api/sirh/paie.xlsx?mois={mois:%Y-%m}", headers=entetes("ADMINRH"))
    feuille = load_workbook(io.BytesIO(r.content)).active
    lignes = list(feuille.iter_rows(values_only=True))
    entete = next(l for l in lignes if l and "Matricule" in l)
    colonne = entete.index("Retenues avances / prêts (DT)")
    ligne = next(l for l in lignes if l and l[0] == "100259")
    assert ligne[colonne] == 100


def test_parametres_reserves_a_l_administrateur(client, entetes):
    corps = {"avance": {"plafond": 1500, "mensualites_max": 4, "taux": 0, "actif": True}}
    assert client.put("/api/prets/types", headers=entetes("100259"), json=corps).status_code == 403
    r = client.put("/api/prets/types", headers=entetes("ADMINRH"), json=corps)
    assert r.status_code == 200 and r.json()["avance"]["plafond"] == 1500
    assert demander(client, entetes, "avance", 1400, 4).status_code == 201
