"""Import Excel RH et tableau de supervision technique."""
import io

from openpyxl import load_workbook

from app.models import Departement, DossierEmploye, Employe, EtatTache, IncidentTechnique, JournalAudit, SoldeConge
from app.services import supervision


def _modele(client, entetes) -> bytes:
    r = client.get("/api/administration/import-rh/modele.xlsx", headers=entetes("ADMINRH"))
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    return r.content


def _ligne(feuille, matricule: str) -> int:
    return next(i for i in range(2, feuille.max_row + 1) if feuille.cell(i, 1).value == matricule)


def _fichier(classeur) -> bytes:
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def test_modele_excel_prempli_et_import_atomique(client, entetes, db):
    autre = Departement(code="AUTRE", nom="Autre direction")
    db.add(autre)
    db.commit()
    classeur = load_workbook(io.BytesIO(_modele(client, entetes)))
    assert {"Instructions", "Références", "Affectations", "Dossiers"} <= set(classeur.sheetnames)
    assert classeur["Références"].sheet_state == "hidden"

    affectations = classeur["Affectations"]
    ligne = _ligne(affectations, "100259")
    affectations.cell(ligne, 4, "AUTRE")
    affectations.cell(ligne, 5, "ADMINRH")
    affectations.cell(ligne, 6, "manager")
    affectations.cell(ligne, 7, "Responsable importé")
    dossiers = classeur["Dossiers"]
    ligne_dossier = _ligne(dossiers, "100259")
    dossiers.cell(ligne_dossier, 4, "15/01/2020")
    dossiers.cell(ligne_dossier, 5, "14/03/1985")
    dossiers.cell(ligne_dossier, 6, "Cadre")
    contenu = _fichier(classeur)

    fichiers = {"fichier": ("import.xlsx", contenu, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    simulation = client.post("/api/administration/import-rh?simulation=true", headers=entetes("ADMINRH"), files=fichiers)
    assert simulation.status_code == 200, simulation.text
    assert simulation.json()["valide"] is True
    assert simulation.json()["collaborateurs_modifies"] == 1
    db.expire_all()
    assert db.query(Employe).filter_by(matricule="100259").one().departement.code == "DIR"

    application = client.post("/api/administration/import-rh?simulation=false", headers=entetes("ADMINRH"), files=fichiers)
    assert application.status_code == 200, application.text
    assert application.json()["appliques"] == 1
    sauvegarde = application.json()["sauvegarde"]
    if db.bind.dialect.name == "sqlite":
        assert sauvegarde.endswith(".db")
    else:
        assert sauvegarde is None
    db.expire_all()
    employe = db.query(Employe).filter_by(matricule="100259").one()
    assert (employe.departement.code, employe.validateur.matricule, employe.niveau, employe.poste) == (
        "AUTRE", "ADMINRH", "manager", "Responsable importé")
    dossier = db.get(DossierEmploye, employe.id)
    assert str(dossier.date_naissance) == "1985-03-14" and dossier.categorie == "Cadre"
    assert db.query(JournalAudit).filter_by(action="import_rh_masse").count() == 1


def test_import_refuse_identite_incoherente_cycle_et_non_rh(client, entetes, db):
    classeur = load_workbook(io.BytesIO(_modele(client, entetes)))
    affectations = classeur["Affectations"]
    ligne = _ligne(affectations, "100130")
    affectations.cell(ligne, 2, "Mauvais nom")
    contenu = _fichier(classeur)
    fichiers = {"fichier": ("import.xlsx", contenu, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post("/api/administration/import-rh?simulation=true", headers=entetes("ADMINRH"), files=fichiers)
    assert r.status_code == 200 and r.json()["valide"] is False
    assert any("identité incohérente" in e for e in r.json()["erreurs"])
    assert client.post("/api/administration/import-rh?simulation=false", headers=entetes("ADMINRH"), files=fichiers).status_code == 422
    assert client.post("/api/administration/import-rh?simulation=true", headers=entetes("100259"), files=fichiers).status_code == 403

    classeur = load_workbook(io.BytesIO(_modele(client, entetes)))
    affectations = classeur["Affectations"]
    affectations.cell(_ligne(affectations, "100130"), 5, "100259")
    contenu = _fichier(classeur)
    r = client.post("/api/administration/import-rh?simulation=true", headers=entetes("ADMINRH"),
                    files={"fichier": ("cycle.xlsx", contenu, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200 and r.json()["valide"] is False
    assert any("boucle" in e for e in r.json()["erreurs"])


def test_supervision_regroupe_erreurs_et_controle_sauvegarde(client, entetes, db):
    assert supervision.executer("test_reussi", lambda: [1, 2]) == [1, 2]

    def echouer():
        raise RuntimeError("panne simulée")

    assert supervision.executer("test_echec", echouer) is None
    assert supervision.executer("test_echec", echouer) is None
    db.expire_all()
    assert db.get(EtatTache, "test_reussi").statut == "succes"
    etat_echec = db.get(EtatTache, "test_echec")
    assert etat_echec.statut == "echec" and etat_echec.echecs == 2
    incident = db.query(IncidentTechnique).filter_by(source="test_echec").one()
    assert incident.occurrences == 2 and incident.resolu_le is None

    assert client.get("/api/supervision", headers=entetes("100259")).status_code == 403
    tableau = client.get("/api/supervision", headers=entetes("ADMINRH"))
    assert tableau.status_code == 200
    assert tableau.json()["incidents_ouverts"] == 1
    assert client.post(f"/api/supervision/incidents/{incident.id}/resoudre", headers=entetes("ADMINRH")).status_code == 200
    db.expire_all()
    assert db.get(IncidentTechnique, incident.id).resolu_le is not None

    # Une sauvegarde forcée existe avant le contrôle d'intégrité.
    assert client.post("/api/sirh/sauvegardes", headers=entetes("ADMINRH")).status_code == 200
    controle = client.post("/api/supervision/sauvegardes/controler", headers=entetes("ADMINRH"))
    assert controle.status_code == 200
    attendu = "ok" if db.bind.dialect.name == "sqlite" else "a_configurer"
    assert controle.json()["statut"] == attendu
    if db.bind.dialect.name == "sqlite":
        assert controle.json()["fichiers"][0]["detail"] == "ok"
    else:
        assert controle.json()["fichiers"] == []


def test_qualite_donnees_signale_les_champs_incomplets_sans_exposer_leurs_valeurs(client, entetes, db):
    employe = db.query(Employe).filter_by(matricule="100259").one()
    employe.departement_id = None
    employe.validateur_id = None
    employe.email = ""
    employe.poste = "Non renseigné"
    employe.date_entree = None
    dossier = db.get(DossierEmploye, employe.id)
    if dossier is None:
        dossier = DossierEmploye(employe_id=employe.id)
        db.add(dossier)
    dossier.date_naissance = None
    dossier.categorie = None
    db.query(SoldeConge).filter_by(employe_id=employe.id).delete()
    db.commit()

    assert client.get("/api/qualite-donnees", headers=entetes("100259")).status_code == 403
    reponse = client.get("/api/qualite-donnees", headers=entetes("ADMINRH"))
    assert reponse.status_code == 200
    rapport = reponse.json()
    erreurs = [a for a in rapport["anomalies"] if a["matricule"] == "100259"]
    assert {a["code"] for a in erreurs} >= {
        "sans_structure", "sans_superieur", "email_absent", "poste_absent",
        "date_entree_absente", "naissance_absente", "categorie_absente", "solde_absent",
    }
    assert all("date_naissance" not in a and "email" not in a for a in erreurs)
